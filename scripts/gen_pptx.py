from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
ESPRESSO=RGBColor(0x1E,0x39,0x32)
CARAMEL=RGBColor(0x00,0x75,0x4A)
GOLD=RGBColor(0xCB,0xA2,0x58)
CREAM=RGBColor(0xF2,0xF0,0xEB)
CREAM_DARK=RGBColor(0xED,0xEB,0xE9)
WHITE=RGBColor(0xFF,0xFF,0xFF)
GREY=RGBColor(0x6B,0x6B,0x6B)
prs=Presentation()
prs.slide_width=Inches(13.33)
prs.slide_height=Inches(7.5)
def bg(slide,color):
    bg=slide.background
    fill=bg.fill
    fill.solid()
    fill.fore_color.rgb=color
def add_shape(slide,left,top,w,h,fill,line=None):
    shape=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,left,top,w,h)
    shape.fill.solid(); shape.fill.fore_color.rgb=fill
    shape.line.fill.background()
    if line: shape.line.color.rgb=line; shape.line.width=Pt(1)
    return shape
def txt(shape,text,size=12,color=ESPRESSO,bold=False,align=PP_ALIGN.LEFT):
    tf=shape.text_frame; tf.word_wrap=True
    p=tf.paragraphs[0]; p.text=text; p.font.size=Pt(size); p.font.color.rgb=color; p.font.bold=bold; p.alignment=align
    return tf
def add_txt(slide,l,t,w,h,text,size=12,color=ESPRESSO,bold=False,align=PP_ALIGN.LEFT):
    tx=slide.shapes.add_textbox(l,t,w,h)
    txt(tx,text,size,color,bold,align)
    return tx
